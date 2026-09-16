from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(r"C:\Users\User\barangayan\barangayan project paper\Barangayan Chapter 3 Research Methodology - Draft.md")
OUTPUT = Path(r"C:\Users\User\barangayan\barangayan project paper\Barangayan Chapter 3 Research Methodology APA 7th Edition.docx")

FONT = "Times New Roman"


def set_font(run, size=12, bold=None, italic=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_markdown_runs(paragraph, text, size=12):
    pattern = re.compile(r"(\*\*.*?\*\*|\*.*?\*)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_font(run, size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, size=size, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, size=size, italic=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_font(run, size=size)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "bottom", "insideH"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
    for edge in ("left", "right", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_font(run)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    header = section.header.paragraphs[0]
    add_page_number(header)
    core = doc.core_properties
    core.title = "Barangayan Research Methodology and Operational Framework"
    core.author = ""
    core.subject = "Chapter 3"
    core.keywords = "Barangayan, research methodology, barangay management system"


def add_body_paragraph(doc, text="", first_line=True, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = alignment
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Inches(0.5) if first_line else Inches(0)
    if text:
        add_markdown_runs(p, text)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Inches(0)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_font(run, bold=True)
    set_keep_with_next(p)
    return p


def add_list_item(doc, text, number=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.left_indent = Inches(0.5)
    pf.first_line_indent = Inches(-0.25)
    prefix = f"{number}. " if number is not None else "- "
    run = p.add_run(prefix)
    set_font(run)
    add_markdown_runs(p, text)
    return p


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.first_line_indent = Inches(0)
    run = p.add_run("\n".join(lines))
    run.font.name = "Courier New"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
    run.font.size = Pt(9)
    return p


def infer_table_title(previous_heading, table_number):
    titles = {
        "3.1.3 Participants and sampling procedure": "Participant Groups and Sampling Plan",
        "3.1.6 Data organization, processing, and analysis": "Interpretation of Weighted Mean Scores",
        "3.2.3 Functional scope": "Functional Scope of Barangayan",
        "3.2.4 User roles and responsibilities": "User Roles and Evaluation Focus",
        "3.3.1 Development approach": "Project Development Phases",
        "3.4.2 Testing levels and criteria": "Testing Levels and Acceptance Criteria",
        "3.4.3 Defect handling and validation of the testing strategy": "Defect Severity and Required Action",
        "3.5.1 Evaluation framework": "Software Quality Areas for Evaluation",
        "3.5.3 Evaluation tasks": "Representative Evaluation Tasks",
        "3.6.1 Work Breakdown Structure": "Work Breakdown Structure",
        "3.7 Computing Standards and Modern Tools and Techniques": "Tools Standards and Techniques",
        "3.8.1 Market model": "Market Model",
        "3.8.2 Measurable benefits": "Expected Benefits and Measures",
        "3.8.3 Three-year product roadmap": "Three Year Product Roadmap",
        "3.9.3 Cost and expense model": "Cost and Expense Model",
        "3.9.4 Success metrics": "Success Metrics",
        "3.9.5 Risks and mitigation": "Risks and Mitigation Measures",
    }
    return titles.get(previous_heading, f"Barangayan Project Information {table_number}")


def add_table(doc, rows, title, number):
    num = doc.add_paragraph()
    num.paragraph_format.space_before = Pt(6)
    num.paragraph_format.space_after = Pt(0)
    num.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    num.paragraph_format.keep_with_next = True
    nrun = num.add_run(f"Table {number}")
    set_font(nrun, bold=True)
    ttl = doc.add_paragraph()
    ttl.paragraph_format.space_before = Pt(0)
    ttl.paragraph_format.space_after = Pt(0)
    ttl.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    ttl.paragraph_format.keep_with_next = True
    trun = ttl.add_run(title)
    set_font(trun, italic=True)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths = [Inches(6.5 / len(rows[0])) for _ in rows[0]]
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.width = widths[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, "F2F2F2")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.first_line_indent = Inches(0)
            run = p.add_run(value)
            set_font(run, size=10, bold=(r_idx == 0))
        if r_idx == 0:
            set_repeat_table_header(table.rows[0])
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(3)
    note.paragraph_format.space_after = Pt(3)
    note.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    note.paragraph_format.first_line_indent = Inches(0)
    return table


def parse_table(lines):
    return [[part.strip() for part in line.strip().strip("|").split("|")] for line in lines if not set(line.replace("|", "").replace("-", "").replace(":", "").replace(" ", ""))]


def is_table_divider(line):
    stripped = line.strip().strip("|").replace("|", "").replace("-", "").replace(":", "").replace(" ", "")
    return stripped == ""


def add_references(doc, lines):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    run = p.add_run("References")
    set_font(run, bold=True)
    for text in lines:
        if not text.strip():
            continue
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.left_indent = Inches(0.5)
        pf.first_line_indent = Inches(-0.5)
        add_markdown_runs(p, text)


def build():
    doc = Document()
    configure_document(doc)

    # APA 7 student title page
    for _ in range(6):
        add_body_paragraph(doc, "", first_line=False, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    title = add_body_paragraph(doc, "Barangayan: Research Methodology and Operational Framework", first_line=False, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    title.runs[0].bold = True
    for _ in range(2):
        add_body_paragraph(doc, "", first_line=False, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    for field in ("[Student Name]", "[Institutional Affiliation]", "[Course Number and Course Name]", "[Instructor Name]", "[Submission Date]"):
        add_body_paragraph(doc, field, first_line=False, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    table_number = 0
    previous_heading = ""
    references = []
    in_references = False
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("## References Relevant to Chapter 3"):
            in_references = True
            i += 1
            continue
        if in_references:
            references.append(line)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("~~~"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("~~~"):
                block.append(lines[i])
                i += 1
            add_code_block(doc, block)
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            rows = []
            for entry in block:
                if not is_table_divider(entry):
                    rows.append([part.strip() for part in entry.strip().strip("|").split("|")])
            if rows:
                table_number += 1
                add_table(doc, rows, infer_table_title(previous_heading, table_number), table_number)
            continue
        if line.startswith("# "):
            add_heading(doc, line[2:].strip(), 1)
            previous_heading = line[2:].strip()
            i += 1
            continue
        if line.startswith("## "):
            add_heading(doc, line[3:].strip(), 1)
            previous_heading = line[3:].strip()
            i += 1
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:].strip(), 2)
            previous_heading = line[4:].strip()
            i += 1
            continue
        if line.startswith("- "):
            add_list_item(doc, line[2:].strip())
            i += 1
            continue
        if len(line) > 3 and line[0].isdigit() and ". " in line[:5]:
            marker, body = line.split(". ", 1)
            if marker.isdigit():
                add_list_item(doc, body.strip(), int(marker))
                i += 1
                continue
        add_body_paragraph(doc, line.strip())
        i += 1

    add_references(doc, references)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
